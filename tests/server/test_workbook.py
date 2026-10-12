"""The headless workbook reader (openpyxl grid path) on a synthetic .xlsx.

Guards the grid logic — sheet/label lookup, the spacer-column value read, the
abbreviated-month ETD, and the VGM container table — without Excel/COM. The .xls
(xlrd) path is validated against live workbooks; here we cover the field logic.
"""

import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))  # repo root
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # tests/ (for sandbox)

import sandbox  # noqa: F401 - keeps the real bootstrap pointer untouched

import openpyxl

from bot_kalung.services.workbook import read_shipment_fields

failures = []


def check(label, cond):
    print(f"{'PASS' if cond else 'FAIL'}  {label}")
    if not cond:
        failures.append(label)


with tempfile.TemporaryDirectory() as tmp:
    folder = Path(tmp) / "16.NIT"
    folder.mkdir()

    wb = openpyxl.Workbook()
    si = wb.active
    si.title = "SI  benar"
    # Destination with a spacer column: label in col 2, value in col 4.
    si.cell(row=5, column=2, value="Port of Discharge")
    si.cell(row=5, column=4, value="KARACHI, PAKISTAN")
    si.cell(row=6, column=2, value="ETD")
    si.cell(row=6, column=3, value="15 AUG 2026")          # abbreviated month
    si.cell(row=7, column=2, value="Party")
    si.cell(row=7, column=3, value="5 X 40'HC")

    vgm = wb.create_sheet("VGM ")
    vgm.cell(row=11, column=2, value="VESSEL NAME")
    vgm.cell(row=11, column=5, value="INTEGRA-184E")
    vgm.cell(row=12, column=2, value="BOOKING NO")
    vgm.cell(row=12, column=5, value="2318229000")
    vgm.cell(row=14, column=4, value="CONTAINER NO")       # table header
    vgm.cell(row=15, column=2, value=1)
    vgm.cell(row=15, column=3, value="40'HC")
    vgm.cell(row=15, column=4, value="cmau8513405")
    vgm.cell(row=16, column=2, value=2)
    vgm.cell(row=16, column=3, value="40'HC")
    vgm.cell(row=16, column=4, value="TRHU5986693")
    wb.save(folder / "NIT16-KATT-VGM,SI,INV,PL.xlsx")

    f = read_shipment_fields(folder)
    check("destination read across a spacer column",
          f.destination_port == "Karachi" and f.destination_country == "Pakistan")
    check("abbreviated-month ETD parsed", f.etd == date(2026, 8, 15))
    check("party size read",
          f.container_quantity == 5 and f.container_size_short == "40'HC")
    check("vessel + voyage read",
          f.vessel_name == "INTEGRA" and f.voyage == "184E")
    check("booking read", f.booking_number == "2318229000")
    check("containers read and upper-cased",
          f.containers == ["CMAU8513405", "TRHU5986693"])
    check("a good workbook reads without warnings", f.warnings == [])

    # HAI's template labels the port "P.O.D", and has an unrelated
    # "... add at destination ..." remark row that must not be taken.
    hai = Path(tmp) / "33.HAI"
    hai.mkdir()
    wb = openpyxl.Workbook()
    si = wb.active
    si.title = "SI "
    si.cell(row=29, column=5, value="P.O.D ")
    si.cell(row=29, column=6, value="NHAVA SHEVA, INDIA")
    si.cell(row=33, column=3, value="Showing agent / add at destination / PIC")
    wb.save(hai / "HAI33-1 iso-Nhava-VGM,SI,Inv,P.List.xlsx")
    f3 = read_shipment_fields(hai)
    check("a 'P.O.D' label is read as the destination (HAI template)",
          f3.destination_port == "Nhava Sheva" and f3.destination_country == "India")

    hai34 = Path(tmp) / "34.HAI"
    hai34.mkdir()
    wb = openpyxl.Workbook()
    si = wb.active
    si.title = "SI"
    si.cell(row=27, column=2, value="SPODUMENE CONCENTRATE")  # contains "POD" — not a label
    si.cell(row=27, column=3, value="WRONG, PLACE")
    si.cell(row=28, column=6, value="POD")
    si.cell(row=28, column=7, value="MAFANG, CHINA")
    wb.save(hai34 / "Hai34-VGM,SI,Inv,P.List.xlsx")
    f4 = read_shipment_fields(hai34)
    check("a bare 'POD' label is read, matched as the whole cell only",
          f4.destination_port == "Mafang" and f4.destination_country == "China")

    # The SI container table's far-right "KET" column (AMI/MFK/NIT/PMA template).
    nit = Path(tmp) / "21.NIT"
    nit.mkdir()
    wb = openpyxl.Workbook()
    si = wb.active
    si.title = "SI"
    for col, head in ((3, "Description of Goods"), (7, "CONTAINER NO"), (8, "SEAL NO"),
                      (11, "tarik"), (12, "KET")):
        si.cell(row=22, column=col, value=head)
    si.cell(row=23, column=7, value="ONEU0968307")
    si.cell(row=23, column=11, value="7/10")
    si.cell(row=23, column=12, value="Tandem 8/10")
    si.cell(row=24, column=7, value=" ffau6222271 ")       # sloppy case/spaces
    si.cell(row=24, column=11, value=date(2026, 10, 8))     # typed as a real date
    si.cell(row=24, column=12, value="Kanopan 8/10 ")
    si.cell(row=25, column=7, value="TLLU5676397")          # no note yet
    si.cell(row=27, column=3, value="TOTAL")
    wb.save(nit / "NIT21-KATT-VGM,SI,INV,PL.xlsx")
    f5 = read_shipment_fields(nit)
    check("the SI 'KET' and 'tarik' columns are read per container",
          f5.container_si == {
              "ONEU0968307": {"note": "Tandem 8/10", "tarik": "7/10"},
              "FFAU6222271": {"note": "Kanopan 8/10", "tarik": "8/10"},  # real date -> d/m
              "TLLU5676397": {"note": "", "tarik": ""}})
    check("a template with no Ket/tarik column leaves them unset",
          f3.container_si is None)

    empty = Path(tmp) / "empty"
    empty.mkdir()
    f2 = read_shipment_fields(empty)
    check("a folder with no workbook warns instead of crashing",
          f2.vessel_name is None and bool(f2.warnings))

print()
if failures:
    print(f"{len(failures)} FAILED: {failures}")
    sys.exit(1)
print("Workbook OK - all checks passed.")
