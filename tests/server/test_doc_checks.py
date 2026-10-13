"""AI-read SI/VGM checks (doc automation 1C step 2): DO vs workbook, BL on-board
date vs the ETD/ATD cells — with a fake reader, no API calls."""

import io
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))  # repo root
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # tests/ (for sandbox)

import sandbox  # noqa: F401 - keeps the real bootstrap pointer untouched

import openpyxl

from bot_kalung.core.db import Database, db_path_for
from bot_kalung.services import doc_checks, si_checks, workbook
from bot_kalung.services.doc_status import DocFile
from bot_kalung.services.excel import ShipmentFields

failures = []


def check(label, cond):
    print(f"{'PASS' if cond else 'FAIL'}  {label}")
    if not cond:
        failures.append(label)


def F(folder, name, created="2026-10-01T00:00:00"):
    return DocFile(folder, name, f"drive:{name}", created, created)


fields = ShipmentFields(booking_number="2335877090", container_quantity=2,
                        container_size_short="40'HC", destination_port="KARACHI",
                        vessel_name="MTT REYA", voyage="26RY149N", etd=date(2026, 10, 20),
                        etd_cells=[("SI", "20 OCTOBER 2026", "2026-10-20"),
                                   ("Inv Buyer", "20 OCTOBER 2026", "2026-10-20")])
DO = {"doc_type": "delivery_order", "number": "2335877090", "booking_number": None,
      "quantity": 2, "size": "40HC", "pod": "Karachi, Pakistan", "vessel": "MTT REYA",
      "voyage": "149N", "etd": "2026-10-20", "on_board_date": None, "issue_date": None}
reads = {}
read = lambda f: reads.get(f.name)          # noqa: E731

files = [F("root", "2335877090.pdf"), F("root", "BATAL-2335734080.pdf"),
         F("root", "Invoice tagihan AMJ28.pdf"), F("root", "IP-KHI-B2A5D0_2026.pdf")]
check("candidates: the booking-named PDF, not Batal / invoices / permits",
      [f.name for f in doc_checks._do_candidates(files, "2335877090")] == ["2335877090.pdf"])

reads["2335877090.pdf"] = DO
r = doc_checks.do_check(files, fields, read)
check("a matching DO is done and names its file",
      r["done"] and r["file"]["name"] == "2335877090.pdf")
reads["2335877090.pdf"] = {**DO, "quantity": 3, "pod": "Kattupalli", "etd": "2026-10-22",
                           "vessel": "MAO GANG", "number": "2339999999"}
r = doc_checks.do_check(files, fields, read)
check("every mismatch is listed", len(r["notes"]) == 4 and not r["done"])
r = doc_checks.do_check(files + [F("kirim", "X-OBL.pdf")], fields, read)
check("...never the DO's ETD (only the BL date counts)",
      len(r["notes"]) == 4 and not any("ETD" in n for n in r["notes"]))
from bot_kalung.services.doc_status import folder_kind
check("HAI/JMI's DO in DO & CC/ is a candidate",
      folder_kind("DO & CC") == "do_cc" and [f.name for f in doc_checks._do_candidates(
          [F("do_cc", "DO SEGMMD26090100.pdf"), F("do_cc", "CC-TASU 1151503.pdf")],
          "SEGMMD26090100")] == ["DO SEGMMD26090100.pdf"])
check("no candidate read yet (no key / budget) keeps the previous check",
      doc_checks.do_check([F("root", "2335877090.pdf")], fields, lambda f: None) is None)
reads["other.pdf"] = {**DO, "doc_type": "other"}
check("candidates that aren't DOs -> DO tidak ditemukan",
      doc_checks.do_check([F("root", "BC other.pdf")], fields,
                          lambda f: {**DO, "doc_type": "other"})["notes"] == ["DO tidak ditemukan"])

check("no BL in Dok kirim -> BL belum diterima",
      doc_checks.atd_check(files, fields, read)["notes"] == ["BL belum diterima"])
bl = [F("kirim", "084600063374-OBL.pdf")]
BL = {**DO, "doc_type": "bill_of_lading", "on_board_date": "2026-10-20"}
check("BL on-board date equal to every ETD cell is done",
      doc_checks.atd_check(bl, fields, lambda f: BL)["done"])
r = doc_checks.atd_check(bl, fields, lambda f: {**BL, "on_board_date": "2026-10-22"})
check("a different on-board date flags each sheet",
      r["notes"] == ["SI: ETD/ATD 20 Oct 2026 ≠ BL 22 Oct 2026",
                     "Inv Buyer: ETD/ATD 20 Oct 2026 ≠ BL 22 Oct 2026"])

check("size normalisation", doc_checks._size("40'HC") == doc_checks._size("40 HIGH CUBE")
      == "40HC" and doc_checks._size("40'") == "40GP")

# ETD/ATD cells from a workbook
wb = openpyxl.Workbook()
wb.active.title = "SI  benar"
wb.active.cell(5, 2, "ETD"), wb.active.cell(5, 4, "20 OCTOBER 2026")
inv = wb.create_sheet("Inv Buyer")
inv.cell(6, 3, "ATD :"), inv.cell(6, 4, "21 OCT 2026")
pl = wb.create_sheet("P.List buyer")
pl.cell(6, 3, "ETD")
buf = io.BytesIO()
wb.save(buf)
cells = si_checks.read_etd_cells(workbook._grids_from_xlsx(buf.getvalue()))
check("ETD / ATD cells are read from SI, INV and PL",
      cells == [("SI  benar", "20 OCTOBER 2026", "2026-10-20"),
                ("Inv Buyer", "21 OCT 2026", "2026-10-21"), ("P.List buyer", None, None)])

# ---- container photos (1D) ------------------------------------------------------
check("ISO 6346 check digit: real numbers pass, a misread fails",
      doc_checks._iso6346_ok("ONEU5311830") and doc_checks._iso6346_ok("CSGU6170181")
      and not doc_checks._iso6346_ok("ONEU5311831"))
check("handwriting tolerance: a misread / dropped digit or seal letters pass",
      doc_checks._handwritten_ok("2335871090", "2335877090")
      and doc_checks._handwritten_ok("233587600", "2335876840")
      and doc_checks._handwritten_ok("1690363", "ID16903AB")
      and not doc_checks._handwritten_ok("2339999999", "2335877090"))

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    for d in ("1.ONEU5311830-galang 8 oct", "2.MISU5170585-kanopan"):
        (root / "Foto" / d).mkdir(parents=True)
        (root / "Foto" / d / "a.jpeg").write_bytes(b"x")
    db = Database(db_path_for(root / "db"))
    db.initialize()
    pf = ShipmentFields(booking_number="MESG11846400",
                        vgm_rows={"ONEU5311830": {"seal": "ID16903AB", "tare": "3,700"},
                                  "MISU5152909": {"seal": "S2", "tare": "3800"}})
    good = {"door_container_number": "ONEU5311830", "tare_kg": 3700,
            "sheet_container_number": "ONEU 5311 830", "sheet_seal_number": "1690363",
            "sheet_do_number": "KESG11846400"}
    doc_checks.new_scan()
    r = doc_checks.photo_check(str(root), pf, db, lambda ref: b"", ask=lambda imgs: good)
    check("photos: a matching container passes; a VGM container without photos and a "
          "photo folder for a container not on the VGM are both named",
          r["notes"] == ["folder foto MISU5170585 tidak ada di VGM", "MISU5152909: foto belum ada"])
    pf.vgm_rows = {"ONEU5311830": pf.vgm_rows["ONEU5311830"]}
    bad = {**good, "tare_kg": 3800, "door_container_number": "ONEU5311831"}
    db.execute("DELETE FROM ai_reads")
    r = doc_checks.photo_check(str(root), pf, db, lambda ref: b"", ask=lambda imgs: bad)
    check("photos: a tare that differs is flagged; a door number failing its check "
          "digit (misread) is not", r["notes"] == ["folder foto MISU5170585 tidak ada di VGM",
                                                  "ONEU5311830: tare foto 3800 ≠ VGM 3700"])

# the cache: one ask per file version, and the per-scan budget
with tempfile.TemporaryDirectory() as tmp:
    db = Database(db_path_for(Path(tmp)))
    db.initialize()
    asked = []
    reader = doc_checks.AiReader(db, lambda ref: b"%PDF", ask=lambda pdf: asked.append(1) or DO)
    doc_checks.new_scan()
    f = F("root", "2335877090.pdf")
    reader(f), reader(f)
    check("a file is asked once, then served from ai_reads", len(asked) == 1)
    reader(F("root", "2335877090.pdf", created="2026-10-02T00:00:00"))
    check("a changed file is read again", len(asked) == 2)
    doc_checks._left = 0
    check("an exhausted budget reads nothing new",
          reader(F("root", "new.pdf")) is None and len(asked) == 2)

print()
if failures:
    print(f"{len(failures)} FAILED: {failures}")
    sys.exit(1)
print("Doc checks OK - all checks passed.")
