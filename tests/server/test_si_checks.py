"""SI/VGM checklist (doc automation 1C step 1) on synthetic workbooks, read
through the same grid loader the worker uses."""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))  # repo root
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # tests/ (for sandbox)

import sandbox  # noqa: F401 - keeps the real bootstrap pointer untouched

import openpyxl

from bot_kalung.services import si_checks, workbook

failures = []


def check(label, cond):
    print(f"{'PASS' if cond else 'FAIL'}  {label}")
    if not cond:
        failures.append(label)


CONTS = [("ONEU0968307", "S1", 3800), ("FFAU6222271", "S2", 3750)]


def build(vgm_rows=CONTS, si_rows=CONTS, inv_rows=CONTS, vgm_date="10 OCTOBER 2026",
          inv_date="10 OCTOBER 2026", si_header="CONTAINER NO", si_gap=False):
    wb = openpyxl.Workbook()
    vgm = wb.active
    vgm.title = "VGM "
    vgm.cell(2, 8, "VGM NO :"), vgm.cell(2, 9, "21102026")
    vgm.cell(3, 8, "DATE :"), vgm.cell(3, 9, vgm_date)
    vgm.cell(5, 5, "MESG11846400")                      # a booking no. — not a container row
    for col, head in ((4, "CONTAINER NO"), (5, "SEAL NO."), (8, "TARE WEIGHT (KG/LBS)")):
        vgm.cell(9, col, head)
    for i, (no, seal, tare) in enumerate(vgm_rows):
        vgm.cell(10 + i, 2, i + 1)                         # numbered rows = expected count
        vgm.cell(10 + i, 4, no), vgm.cell(10 + i, 5, seal), vgm.cell(10 + i, 8, tare)
    si = wb.create_sheet("SI  benar")
    si.cell(17, 7, si_header), si.cell(17, 8, "SEAL NO.")
    for i, (no, seal, _) in enumerate(si_rows):
        r = 19 + i * (2 if si_gap else 1)                 # AMJ: HS-code lines between rows
        si.cell(r, 7, no), si.cell(r, 8, seal)
    inv = wb.create_sheet("Inv Buyer")
    inv.cell(3, 5, "DATE"), inv.cell(3, 6, inv_date)
    inv.cell(11, 5, "CONTAINER NO"), inv.cell(11, 6, "SEAL NO")
    for i, (no, seal, _) in enumerate(inv_rows):
        inv.cell(12 + i, 5, no), inv.cell(12 + i, 6, seal)
    buf = io.BytesIO()
    wb.save(buf)
    return workbook._grids_from_xlsx(buf.getvalue())


def run(**kw):
    return {c["key"]: c for c in si_checks.compute(build(**kw), 2)}


ok = run()
check("a complete workbook ticks both checks",
      ok["containers"]["done"] and ok["dates"]["done"])
check("the booking number on the VGM is not mistaken for a container",
      ok["containers"]["notes"] == [])
check("AMJ's 'CONT NO.' header with HS-code lines between rows still reads",
      run(si_header="CONT NO.", si_gap=True)["containers"]["done"])

r = run(inv_rows=CONTS[:1])
check("a container missing from INV is reported", not r["containers"]["done"]
      and "Inv Buyer: 1/2 kontainer terisi" in r["containers"]["notes"])
r = run(si_rows=[("ONEU0968307", "", 0), ("FFAU6222271", "S2", 0)])
check("an empty seal is reported", "SI  benar: seal kosong (1)" in r["containers"]["notes"])
r = run(vgm_rows=[("ONEU0968307", "S1", None), ("FFAU6222271", "S2", 3750)])
check("an empty VGM tare is reported", "VGM: tare kosong (1)" in r["containers"]["notes"])
r = run(inv_rows=[("ONEU0968307", "S1", 0), ("TCNU4466857", "S9", 0)])
check("a container on INV but not on the VGM is reported",
      "Inv Buyer: TCNU4466857 tidak ada di VGM" in r["containers"]["notes"])
r = run(vgm_rows=[(None, None, None), (None, None, None)], si_rows=[], inv_rows=[])
check("an unfilled workbook shows 0/N on every sheet",
      "VGM: 0/2 kontainer terisi" in r["containers"]["notes"]
      and "Inv Buyer: 0/2 kontainer terisi" in r["containers"]["notes"])

r = run(vgm_date="OCTOBER 2026")
check("a date without a day is not complete",
      not r["dates"]["done"] and r["dates"]["notes"] == ["VGM: tanggal belum lengkap (OCTOBER 2026)"])
check("both VGM and INV dates are checked",
      run(inv_date="NOVEMBER 2026")["dates"]["notes"] == ["Inv Buyer: tanggal belum lengkap (NOVEMBER 2026)"])

# ---- document numbers -------------------------------------------------------
def numbers(vgm_cell, inv=None, pl=None, own=None):
    """vgm_cell: ("inline", "NO : VGM-30102026") or ("label", "30102026")."""
    wb = openpyxl.Workbook()
    vgm = wb.active
    vgm.title = "VGM "
    if vgm_cell[0] == "inline":
        vgm.cell(3, 2, vgm_cell[1])
    else:
        vgm.cell(2, 7, "VGM NO :"), vgm.cell(2, 8, vgm_cell[1])
    for i, v in enumerate(inv if isinstance(inv, list) else ([inv] if inv is not None else [])):
        sh = wb.create_sheet(f"Inv {i}")
        sh.cell(2, 5, "INVOICE NO :"), sh.cell(2, 6, v)
        sh.cell(3, 5, "DATE"), sh.cell(3, 6, "25 NOVEMBER 2026")   # month never compared
    if pl is not None:
        sh = wb.create_sheet("P.List buyer")
        sh.cell(2, 5, "PACKING LIST NO :"), sh.cell(2, 6, pl)
    buf = io.BytesIO()
    wb.save(buf)
    raw = si_checks.read_numbers(workbook._grids_from_xlsx(buf.getvalue()))
    return si_checks.numbers_check(raw, own or set())


check("prefixed VGM -> I- / PL- numbers pass",
      numbers(("inline", "NO : VGM-30102026"), "I-30102026", "PL-30102026", {30})["done"])
check("plain numbers pass on every document (the invoice date's month is not compared)",
      numbers(("label", "33102026"), ["33102026", "33102026"], "33102026", {33})["done"])
r = numbers(("inline", "NO : VGM-02092026"), "02092026", "02092026", {2})
check("a prefixed VGM with plain INV/PL is flagged with the expected number",
      r["notes"] == ["Inv 0: 02092026 (seharusnya I-02092026)",
                     "P.List buyer: 02092026 (seharusnya PL-02092026)"])
check("a seq that isn't this shipment's (copied workbook) is flagged",
      "VGM: 37102026 — nomor urut 37 bukan untuk pengiriman ini"
      in numbers(("label", "37102026"), "37102026", "37102026", {40})["notes"])
r = numbers(("label", "33102026"), ["33102026", "33092026"], None, {33})
check("documents that disagree with the VGM are flagged",
      r["notes"] == ["Inv 1: 33092026 — tidak sama dengan VGM (33102026)"])
check("JMI's own pattern passes when consistent",
      numbers(("inline", "NO : VGM-27092026"), "27/JMI/092026", "PL/JMI/092026/27", {27})["done"])
check("AMI's pattern (Roman month, 2-digit year, workbook code number) passes",
      numbers(("inline", "NO : VGM-031092026"), "031/AMI/IX/26", "031/AMI/IX/26", {6, 31})["done"])
check("AMI fully in its own pattern (VGM too) passes",
      numbers(("label", "034/AMI/X/26"), "034/AMI/X/26", "034/AMI/X/26", {7, 34})["done"])
check("CKJ's pattern passes", numbers(("inline", "NO : VGM-24082026"), "24/INV/CKJ/VIII/2026",
                                      None, {24})["done"])
check("an exporter pattern that disagrees is still caught",
      not numbers(("inline", "NO : VGM-27092026"), "26/JMI/092026", None, {27})["done"])
check("a missing INV number is reported",
      numbers(("label", "33102026"), "", None, {33})["notes"] == ["Inv 0: nomor kosong"])

# ---- AMI/PMA follow the buyer's ordering: seq = code number in the folder name
import json
import tempfile
from bot_kalung.core.db import Database, db_path_for, new_id
from bot_kalung.services import tracker
from bot_kalung.services.action_items import ActionItems
from bot_kalung.services.excel import ShipmentFields

with tempfile.TemporaryDirectory() as tmp:
    db = Database(db_path_for(Path(tmp)))
    db.initialize()
    sid = new_id()
    db.execute("INSERT INTO shipments (id, exporter_code, sequence_number, status, created_at) "
               "VALUES (?, 'PMA', 1, 'active', '2026-09-21')", (sid,))
    ActionItems(db).seed(sid, "PMA", "India")

    def pma_numbers(vgm, inv):
        f = ShipmentFields()
        f.workbook = "PMA02-Katt-VGM,SI,INV,PL.xlsx"
        f.doc_numbers = {"docs": [("VGM", "vgm", vgm), ("INV Buyer-BC", "inv", inv)]}
        tracker._store_si_checks(db, sid, "PMA", f, 1, "1.PMA02-4x40-Katt-T22937-Reya142N")
        return json.loads(db.query_one("SELECT checks FROM action_items WHERE shipment_id=? "
                                       "AND code='si_vgm'", (sid,))["checks"])

    got = pma_numbers("VGM-002092026", "002/PMA/IX/26")
    check("PMA folder '1.PMA02': documents numbered 2 pass, and only the number check is stored",
          [c["key"] for c in got] == ["numbers"] and got[0]["done"])
    check("...documents numbered 1 (our folder order) are flagged",
          not pma_numbers("VGM-001092026", "001/PMA/IX/26")[0]["done"])

    for sid2, fname, want in (("h40", "40.CP26167B-1 iso-Nhava-TMEL-GAN-WH101-N388", 1),
                              ("h41", "41.CP26200-Nhava-Isolasi-WH101", 0)):
        db.execute("INSERT INTO shipments (id, exporter_code, sequence_number, status, created_at) "
                   "VALUES (?, 'HAI', ?, 'active', '2026-10-01')", (sid2, int(sid2[1:])))
        tracker._store_si_checks(db, sid2, "HAI", ShipmentFields(), int(sid2[1:]), fname)
        check(f"HAI folder {fname[:24]!r}: ISO tag = {want} ('iso' as a word only)",
              db.query_one("SELECT iso FROM shipments WHERE id=?", (sid2,))["iso"] == want)

empty = openpyxl.Workbook()
buf = io.BytesIO()
empty.save(buf)
r = {c["key"]: c for c in si_checks.compute(workbook._grids_from_xlsx(buf.getvalue()), 2)}
check("no container table at all is never 'done'",
      not r["containers"]["done"] and "tabel kontainer tidak ditemukan" in r["containers"]["notes"])

print()
if failures:
    print(f"{len(failures)} FAILED: {failures}")
    sys.exit(1)
print("SI checks OK - all checks passed.")
